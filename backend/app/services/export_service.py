"""Exporting a user's posts as a PDF or Word document.

Both formats share the same two steps: load the posts, then walk each post's Markdown
content with `app.core.markdown` and add one block at a time to the document.
"""

import io
from pathlib import Path
from typing import Literal

from docx import Document
from docx.document import Document as DocxDocument
from docx.shared import Pt
from docx.text.paragraph import Paragraph as DocxParagraph
from fpdf import FPDF
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker

from app.core.markdown import Block, CodeBlock, Heading, ListBlock, Paragraph, Run, parse_blocks
from app.models import Post, PostStatus, User
from app.repositories.post_repository import PostRepository

ExportFormat = Literal["pdf", "docx"]

FONT_DIR = Path(__file__).resolve().parent.parent / "assets" / "fonts"

# A document is built whole in memory, unlike the row-at-a-time CSV export this replaced,
# so a very prolific author can't make the server buffer an unbounded file
MAX_EXPORT_POSTS = 500


async def _load_posts(
    sessionmaker: async_sessionmaker[AsyncSession],
    user: User,
    status: PostStatus | None,
    search: str | None,
) -> list[Post]:
    """Up to MAX_EXPORT_POSTS of the user's posts, oldest first.

    Opens its own database session: the caller builds the document after this
    returns, once the request's own session may already have closed.
    """
    async with sessionmaker() as session:
        posts = []
        async for post in PostRepository(session).stream_by_author(user.id, status, search):
            posts.append(post)
            if len(posts) >= MAX_EXPORT_POSTS:
                break
        # Touches every attribute the document builders read, so it is all loaded
        # before the session above closes and the posts become detached
        for post in posts:
            _ = (post.author.username, post.like_count, post.comment_count)
            _ = [topic.name for topic in post.topics]
        return posts


def _meta_line(post: Post) -> str:
    if post.status is PostStatus.PUBLISHED and post.published_at:
        return f"Published {post.published_at:%B %d, %Y}"
    return f"Draft, last edited {post.updated_at:%B %d, %Y}"


def _topics_line(post: Post) -> str | None:
    if not post.topics:
        return None
    return "Topics: " + ", ".join(topic.name for topic in post.topics)


# --- PDF (fpdf2, with a bundled Unicode font so “smart quotes”, dashes and emoji
# never crash the built-in Latin-1-only core fonts) ---


def _pdf_document() -> FPDF:
    pdf = FPDF()
    pdf.set_margins(20, 20, 20)
    pdf.set_auto_page_break(True, margin=20)
    for style, suffix in [("", ""), ("B", "-Bold"), ("I", "-Oblique"), ("BI", "-BoldOblique")]:
        pdf.add_font("DejaVu", style, str(FONT_DIR / f"DejaVuSans{suffix}.ttf"))
    pdf.add_font("DejaVuMono", "", str(FONT_DIR / "DejaVuSansMono.ttf"))
    pdf.add_page()
    return pdf


def _pdf_runs(pdf: FPDF, runs: list[Run], size: float) -> None:
    for run in runs:
        if run.code:
            pdf.set_font("DejaVuMono", "", size - 1)
        else:
            pdf.set_font("DejaVu", ("B" if run.bold else "") + ("I" if run.italic else ""), size)
        for line in run.text.split("\n"):
            if line:
                pdf.write(h=size * 0.55, text=line)
    pdf.ln(size * 0.8)


def _pdf_block(pdf: FPDF, block: Block) -> None:
    if isinstance(block, Heading):
        _pdf_runs(pdf, block.runs, max(11, 18 - 2 * block.level))
    elif isinstance(block, Paragraph):
        _pdf_runs(pdf, block.runs, 11)
    elif isinstance(block, ListBlock):
        for i, item in enumerate(block.items):
            bullet = f"{i + 1}." if block.ordered else "-"
            pdf.set_font("DejaVu", "", 11)
            pdf.write(h=6, text=f"{bullet} ")
            _pdf_runs(pdf, item, 11)
    elif isinstance(block, CodeBlock):
        pdf.set_font("DejaVuMono", "", 9.5)
        pdf.multi_cell(w=0, text=block.text, border=1, padding=3)
        pdf.ln(2)


def _pdf_post(pdf: FPDF, post: Post) -> None:
    pdf.set_font("DejaVu", "B", 16)
    pdf.multi_cell(w=0, text=post.title, new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("DejaVu", "I", 10)
    lines = [_meta_line(post)]
    if topics := _topics_line(post):
        lines.append(topics)
    pdf.multi_cell(w=0, text=" · ".join(lines), new_x="LMARGIN", new_y="NEXT")
    pdf.ln(4)
    for block in parse_blocks(post.content):
        _pdf_block(pdf, block)


def _build_pdf(posts: list[Post]) -> bytes:
    pdf = _pdf_document()
    if not posts:
        pdf.set_font("DejaVu", "I", 12)
        pdf.write(text="No posts to export.")
    for i, post in enumerate(posts):
        if i > 0:
            pdf.add_page()
        _pdf_post(pdf, post)
    return bytes(pdf.output())


# --- Word (python-docx) ---


def _docx_runs(paragraph: DocxParagraph, runs: list[Run]) -> None:
    for run in runs:
        text_run = paragraph.add_run(run.text)
        text_run.bold = run.bold
        text_run.italic = run.italic
        if run.code:
            text_run.font.name = "Courier New"


def _docx_block(document: DocxDocument, block: Block) -> None:
    if isinstance(block, Heading):
        _docx_runs(document.add_heading(level=min(block.level, 9)), block.runs)
    elif isinstance(block, Paragraph):
        _docx_runs(document.add_paragraph(), block.runs)
    elif isinstance(block, ListBlock):
        style = "List Number" if block.ordered else "List Bullet"
        for item in block.items:
            _docx_runs(document.add_paragraph(style=style), item)
    elif isinstance(block, CodeBlock):
        code_run = document.add_paragraph().add_run(block.text)
        code_run.font.name = "Courier New"
        code_run.font.size = Pt(10)


def _docx_post(document: DocxDocument, post: Post) -> None:
    document.add_heading(post.title, level=1)
    lines = [_meta_line(post)]
    if topics := _topics_line(post):
        lines.append(topics)
    meta = document.add_paragraph()
    meta.add_run(" · ".join(lines)).italic = True
    for block in parse_blocks(post.content):
        _docx_block(document, block)


def _build_docx(posts: list[Post]) -> bytes:
    document = Document()
    if not posts:
        document.add_paragraph("No posts to export.")
    for i, post in enumerate(posts):
        if i > 0:
            document.add_page_break()
        _docx_post(document, post)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


async def export_posts(
    sessionmaker: async_sessionmaker[AsyncSession],
    user: User,
    export_format: ExportFormat,
    status: PostStatus | None,
    search: str | None,
) -> bytes:
    """The user's posts, drafts included, as a downloadable PDF or Word document."""
    posts = await _load_posts(sessionmaker, user, status, search)
    return _build_pdf(posts) if export_format == "pdf" else _build_docx(posts)
